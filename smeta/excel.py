"""Smetani .xlsx faylga yozish (openpyxl).

Hisob-kitob brauzerda (app.js) bajariladi, shuning uchun ilova tayyor "varaq modeli"ni yuboradi:
ekrandagi "Excel ko'rinishi" ham, shu fayl ham bitta modeldan quriladi — ikkalasi bir xil.

Model:
{"sheets": [{"name": "Smeta", "cols": [5, 36, ...], "freeze": 4, "table": [3, 40],
             "rows": [[{"v": "matn yoki son", "f": "money|dec2|int", "s": "title|meta|head|group|sub|total|grand|b"}, null, ...], ...]}]}
"""
import math
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MAX_SHEETS, MAX_ROWS, MAX_COLS, MAX_TEXT = 5, 20000, 40, 2000

FORMATS = {"money": "#,##0", "dec2": "0.00", "int": "0"}


def _fill(color):
    return PatternFill("solid", fgColor=color)


# Ranglar ilova palitrasi bilan bir xil (yashil + qora)
STYLES = {
    "title": {"font": Font(bold=True, size=14, color="0C140F")},
    "meta": {"font": Font(color="536159")},
    "head": {"font": Font(bold=True, color="FFFFFF"), "fill": _fill("15803D"),
             "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True)},
    "group": {"font": Font(bold=True, color="0C140F"), "fill": _fill("DCFCE7")},
    "sub": {"font": Font(bold=True), "fill": _fill("EEF3EF")},
    "total": {"font": Font(bold=True)},
    "grand": {"font": Font(bold=True, size=12, color="4ADE80"), "fill": _fill("0C140F")},
    "b": {"font": Font(bold=True)},
}
LINE = Side(style="thin", color="C9D5CC")
BORDER = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)


def _sheet_name(name, used):
    name = re.sub(r"[\[\]:*?/\\]", " ", str(name or "Smeta")).strip()[:31] or "Smeta"
    base, i = name, 2
    while name.lower() in used:
        name = f"{base[:28]} {i}"
        i += 1
    used.add(name.lower())
    return name


def _value(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        raise ValueError("mantiqiy qiymat kutilmagan")
    if isinstance(v, (int, float)):
        if not math.isfinite(v):
            raise ValueError("son noto'g'ri")
        return v
    if isinstance(v, str):
        return v[:MAX_TEXT]
    raise ValueError("katak qiymati noto'g'ri")


def build_workbook(data):
    """Varaq modelidan .xlsx baytlarini qaytaradi. Noto'g'ri ma'lumotda ValueError."""
    if not isinstance(data, dict) or not isinstance(data.get("sheets"), list):
        raise ValueError("sheets yo'q")
    sheets = data["sheets"]
    if not 1 <= len(sheets) <= MAX_SHEETS:
        raise ValueError("varaqlar soni noto'g'ri")

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "SmetaGo"
    used = set()
    for sh in sheets:
        if not isinstance(sh, dict):
            raise ValueError("varaq noto'g'ri")
        rows, cols = sh.get("rows") or [], sh.get("cols") or []
        if not isinstance(rows, list) or len(rows) > MAX_ROWS or not isinstance(cols, list) or len(cols) > MAX_COLS:
            raise ValueError("varaq hajmi noto'g'ri")
        ws = wb.create_sheet(_sheet_name(sh.get("name"), used))

        for ci, w in enumerate(cols, 1):
            if isinstance(w, (int, float)) and 1 <= w <= 120:
                ws.column_dimensions[ws.cell(1, ci).column_letter].width = w

        table = sh.get("table")
        t_from, t_to = (table if isinstance(table, list) and len(table) == 2
                        and all(isinstance(x, int) for x in table) else (-1, -2))

        for ri, row in enumerate(rows, 1):
            if row is None:
                continue
            if not isinstance(row, list) or len(row) > MAX_COLS:
                raise ValueError("qator noto'g'ri")
            in_table = t_from <= ri - 1 <= t_to
            for ci in range(1, max(len(row), len(cols) if in_table else 0) + 1):
                cell_def = row[ci - 1] if ci <= len(row) else None
                if cell_def is not None and not isinstance(cell_def, dict):
                    raise ValueError("katak noto'g'ri")
                cell_def = cell_def or {}
                value = _value(cell_def.get("v"))
                style = STYLES.get(cell_def.get("s"))
                if value is None and not style and not in_table:
                    continue
                cell = ws.cell(ri, ci)
                if value is not None:
                    cell.value = value
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"  # foydalanuvchi matni formula bo'lib bajarilmasin
                    fmt = FORMATS.get(cell_def.get("f"))
                    if fmt and isinstance(value, (int, float)):
                        cell.number_format = fmt
                if style:
                    for attr, val in style.items():
                        setattr(cell, attr, val)
                if in_table:
                    cell.border = BORDER
                    if isinstance(value, str) and cell_def.get("s") != "head":
                        cell.alignment = Alignment(vertical="top")
            if in_table and any(isinstance(c, dict) and c.get("s") == "head" for c in row):
                ws.row_dimensions[ri].height = 32

        freeze = sh.get("freeze")
        if isinstance(freeze, int) and 0 < freeze < len(rows):
            ws.freeze_panes = f"A{freeze + 1}"
            ws.print_title_rows = f"{freeze}:{freeze}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
