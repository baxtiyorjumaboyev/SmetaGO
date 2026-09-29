import json

from django.contrib.auth.models import User
from django.contrib.staticfiles import finders
from django.core.cache import cache
from django.test import Client, TestCase
from django.urls import reverse

from .malumotnoma import SEED_FILE, build_reference
from .models import CatalogItem, Material, Obyekt

STATE = {"v": 1, "obj": {"name": "Chilonzor kvartira", "region": "Toshkent sh."}, "rooms": [{}, {}]}


class SmetaTests(TestCase):
    def setUp(self):
        cache.clear()  # ma'lumotnoma keshi testlar orasida o'tib ketmasin
        self.user = User.objects.create_user("ali", password="SmetaGo-2026!")
        self.c = Client()
        self.c.force_login(self.user)

    def test_register_and_login_pages(self):
        anon = Client()
        landing = anon.get("/")  # mehmonga — sayt (bosh sahifa)
        self.assertEqual(landing.status_code, 200)
        self.assertContains(landing, reverse("register"))
        self.assertContains(landing, 'rel="manifest"')
        self.assertRedirects(anon.get(reverse("obyekt_create")), "/kirish/?next=/obyekt/yangi/")
        # "Parolni unutdingizmi?" — faqat xat yuborish sozlangan bo'lsa
        self.assertNotContains(anon.get(reverse("login")), reverse("password_reset"))
        with self.settings(EMAIL_HOST="smtp.example.com"):
            self.assertContains(anon.get(reverse("login")), reverse("password_reset"))
        form = {"username": "vali", "password1": "Qurilish-2026!", "password2": "Qurilish-2026!"}
        self.assertNotContains(anon.get(reverse("register")), 'name="email"')  # email so'ralmaydi
        r = anon.post(reverse("register"), form)
        self.assertRedirects(r, "/")
        self.assertEqual(User.objects.get(username="vali").email, "")
        # shu login bilan ikkinchi hisob ochilmaydi
        r = Client().post(reverse("register"), form)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(User.objects.filter(username="vali").count(), 1)

    def test_password_reset(self):
        from django.core import mail

        self.user.email = "ali@example.com"
        self.user.save()
        anon = Client()
        # noma'lum email: xat ketmaydi, lekin sahifa bir xil (hisob borligi oshkor bo'lmaydi)
        r = anon.post(reverse("password_reset"), {"email": "yoq@example.com"})
        self.assertRedirects(r, reverse("password_reset_done"))
        self.assertEqual(len(mail.outbox), 0)
        r = anon.post(reverse("password_reset"), {"email": "ALI@example.com"})
        self.assertRedirects(r, reverse("password_reset_done"))
        self.assertEqual(len(mail.outbox), 1)
        msg = mail.outbox[0]
        self.assertEqual(msg.to, ["ali@example.com"])
        self.assertIn("ali", msg.body)
        link = next(w for w in msg.body.split() if "/parol-tiklash/" in w)
        path = link.split("://", 1)[1].split("/", 1)[1]
        r = anon.get("/" + path, follow=True)  # token sessiyaga olinadi, forma ochiladi
        self.assertContains(r, 'name="new_password1"')
        r = anon.post(r.redirect_chain[-1][0], {"new_password1": "Yangi-parol-2026", "new_password2": "Yangi-parol-2026"})
        self.assertRedirects(r, reverse("password_reset_complete"))
        self.assertTrue(Client().login(username="ali", password="Yangi-parol-2026"))
        # havola bir marta ishlaydi
        self.assertContains(Client().get("/" + path, follow=True), reverse("password_reset"))
        self.assertNotContains(Client().get("/" + path, follow=True), 'name="new_password1"')

    def test_create_open_and_save(self):
        r = self.c.post(reverse("obyekt_create"))
        o = Obyekt.objects.get(owner=self.user)
        self.assertRedirects(r, reverse("obyekt_app", args=[o.pk]))
        page = self.c.get(reverse("obyekt_app", args=[o.pk])).content.decode()
        self.assertIn('id="smeta-state"', page)
        self.assertIn("window.SMETAGO", page)
        self.assertIn("/static/smeta/js/app.js", page)

        url = reverse("obyekt_state", args=[o.pk])
        r = self.c.put(url, json.dumps(STATE), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        o.refresh_from_db()
        self.assertEqual(o.name, "Chilonzor kvartira")
        self.assertEqual(o.rooms_count, 2)
        self.assertEqual(self.c.get(url).json(), STATE)

        list_page = self.c.get("/").content.decode()
        self.assertIn("Chilonzor kvartira", list_page)
        self.assertIn("2 xona", list_page)

    def test_sample_object(self):
        self.c.post(reverse("obyekt_create"), {"namuna": "1"})
        self.assertEqual(Obyekt.objects.get().state, {"namuna": True})

    def test_bad_state_rejected(self):
        o = Obyekt.objects.create(owner=self.user)
        url = reverse("obyekt_state", args=[o.pk])
        self.assertEqual(self.c.put(url, "{bad", content_type="application/json").status_code, 400)
        self.assertEqual(self.c.put(url, '{"v":2}', content_type="application/json").status_code, 400)
        self.assertEqual(self.c.post(url).status_code, 405)

    def test_copy_and_delete(self):
        o = Obyekt.objects.create(owner=self.user, name="A", state=dict(STATE, obj={"name": "A"}))
        self.c.post(reverse("obyekt_copy", args=[o.pk]))
        copy = Obyekt.objects.exclude(pk=o.pk).get()
        self.assertEqual(copy.name, "A (nusxa)")
        self.assertEqual(copy.state["obj"]["name"], "A (nusxa)")
        self.assertEqual(Obyekt.objects.get(pk=o.pk).state["obj"]["name"], "A")
        self.c.post(reverse("obyekt_delete", args=[o.pk]))
        self.assertFalse(Obyekt.objects.filter(pk=o.pk).exists())

    def test_other_users_objects_hidden(self):
        other = User.objects.create_user("begona", password="x")
        o = Obyekt.objects.create(owner=other, state=STATE)
        self.assertEqual(self.c.get(reverse("obyekt_app", args=[o.pk])).status_code, 404)
        self.assertEqual(self.c.get(reverse("obyekt_state", args=[o.pk])).status_code, 404)
        r = self.c.put(reverse("obyekt_state", args=[o.pk]), json.dumps(STATE), content_type="application/json")
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.c.post(reverse("obyekt_delete", args=[o.pk])).status_code, 404)

    def test_reference_matches_seed(self):
        """Bazadan yig'ilgan ma'lumotnoma asl data.js (seed) bilan bir xil bo'lishi kerak."""
        seed = json.loads(SEED_FILE.read_text(encoding="utf-8"))
        ref = build_reference()
        self.assertEqual(ref["catalog"], seed["catalog"])
        self.assertEqual(ref["prices"], seed["prices"])
        self.assertEqual(list(ref["roomTypes"]), list(seed["room_types"]))
        for name, rt in seed["room_types"].items():
            got = ref["roomTypes"][name]
            self.assertEqual((got["floor"], got["wall"], got["ceil"]), (rt["floor"], rt["wall"], rt["ceil"]))
            self.assertEqual(sorted(got["s"]), sorted(rt["s"]))

    def test_admin_changes_reach_app_page(self):
        o = Obyekt.objects.create(owner=self.user)
        self.c.get(reverse("obyekt_app", args=[o.pk]))  # ma'lumotnoma keshga tushadi
        # admin kabi save() orqali o'zgartirish — kesh signal bilan tozalanishi kerak
        for key, field, value in (("divan", "price", 5100000), ("seyf", "active", False)):
            item = CatalogItem.objects.get(key=key)
            setattr(item, field, value)
            item.save()
        lam = Material.objects.get(key="laminat")
        lam.src1 = 99000
        lam.save()
        page = self.c.get(reverse("obyekt_app", args=[o.pk])).content.decode()
        ref = json.loads(page.split('id="smeta-ref" type="application/json">')[1].split("</script>")[0])
        items = {it["id"]: it for g in ref["catalog"] for it in g["items"]}
        self.assertEqual(items["divan"]["p"], 5100000)
        self.assertNotIn("seyf", items)
        self.assertNotIn("seyf", ref["roomTypes"]["Ofis xonasi"]["s"])
        self.assertEqual(next(p for p in ref["prices"] if p["id"] == "laminat")["src"][0], 99000)

    def test_reference_admin_pages(self):
        admin = User.objects.create_superuser("admin", password="x")
        c = Client()
        c.force_login(admin)
        for name in ("cataloggroup", "catalogitem", "material", "roomtype", "obyekt"):
            self.assertEqual(c.get(f"/admin/smeta/{name}/").status_code, 200, name)
        item = CatalogItem.objects.get(key="shkaf")
        self.assertEqual(c.get(f"/admin/smeta/catalogitem/{item.pk}/change/").status_code, 200)

    def test_default_language_is_uzbek(self):
        r = self.c.get("/", HTTP_ACCEPT_LANGUAGE="ru")  # brauzer tili hisobga olinmaydi
        self.assertContains(r, '<html lang="uz">')
        self.assertContains(r, "Obyektlar")

    def test_switch_to_russian(self):
        r = self.c.post("/i18n/setlang/", {"language": "ru", "next": "/"})
        self.assertRedirects(r, "/", fetch_redirect_response=False)
        page = self.c.get("/")
        self.assertContains(page, '<html lang="ru">')
        self.assertContains(page, "Объекты")
        self.assertContains(page, "+ Новый объект")
        self.c.post(reverse("obyekt_create"))
        self.assertEqual(Obyekt.objects.get().name, "Новый объект")

    def test_app_page_reference_in_russian(self):
        self.c.cookies["django_language"] = "ru"
        o = Obyekt.objects.create(owner=self.user)
        page = self.c.get(reverse("obyekt_app", args=[o.pk])).content.decode()
        self.assertIn('<html lang="ru">', page)
        self.assertIn("/static/smeta/js/i18n.js", page)
        ref = json.loads(page.split('id="smeta-ref" type="application/json">')[1].split("</script>")[0])
        self.assertEqual(ref["catalog"][0]["g"], "Мебель")
        items = {it["id"]: it for g in ref["catalog"] for it in g["items"]}
        self.assertEqual(items["shkaf"]["n"], "Шкаф")
        self.assertEqual(items["shkaf"]["v"][0][0], "2-дверный")
        # xona turi kaliti o'zbekcha qoladi (obyektlarda saqlanadi), nomi — ruscha
        self.assertEqual(ref["roomTypes"]["Mehmonxona"]["l"], "Гостиная")
        lam = next(p for p in ref["prices"] if p["id"] == "laminat")
        self.assertEqual((lam["n"], lam["g"]), ("Ламинат", "Pol"))

    def test_empty_russian_name_falls_back_to_uzbek(self):
        CatalogItem.objects.filter(key="divan").update(name_ru="")
        items = {it["id"]: it for g in build_reference("ru")["catalog"] for it in g["items"]}
        self.assertEqual(items["divan"]["n"], "Divan")

    def test_pwa_manifest(self):
        r = Client().get("/manifest.webmanifest")
        self.assertEqual(r["Content-Type"], "application/manifest+json")
        m = r.json()
        self.assertEqual((m["display"], m["scope"], m["short_name"]), ("standalone", "/", "SmetaGo"))
        sizes = {(i["sizes"], i["purpose"]) for i in m["icons"]}
        self.assertTrue({("192x192", "any"), ("512x512", "any"), ("512x512", "maskable")} <= sizes)
        for icon in m["icons"]:  # ikonka fayllari haqiqatan bor
            self.assertTrue(finders.find(icon["src"].removeprefix("/static/")), icon["src"])
        c = Client()
        c.cookies["django_language"] = "ru"
        self.assertEqual(c.get("/manifest.webmanifest").json()["lang"], "ru")

    def test_service_worker(self):
        r = Client().get("/sw.js")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r["Content-Type"].startswith("application/javascript"))
        self.assertIn("no-cache", r["Cache-Control"])
        js = r.content.decode()
        self.assertIn('"/offline/"', js)
        self.assertIn("/static/smeta/js/app.js", js)
        precache = json.loads(js.split("const PRECACHE = ")[1].split(";")[0])
        for url in precache:  # oldindan keshlanadigan hamma fayl mavjud bo'lishi shart
            if url.startswith("/static/"):
                self.assertTrue(finders.find(url.removeprefix("/static/")), url)
        self.assertEqual(Client().get("/offline/").status_code, 200)

    def test_app_page_is_installable(self):
        o = Obyekt.objects.create(owner=self.user)
        page = self.c.get(reverse("obyekt_app", args=[o.pk])).content.decode()
        for needle in ('rel="manifest"', 'name="theme-color"', "apple-touch-icon", "smeta/js/pwa.js", "updated:"):
            self.assertIn(needle, page)

    XLSX_SHEETS = {"sheets": [
        {"name": "Smeta", "cols": [5, 30, 12], "freeze": 2, "table": [1, 3], "rows": [
            [{"v": "SMETA: Test", "s": "title"}],
            [{"v": "№", "s": "head"}, {"v": "Nomi", "s": "head"}, {"v": "Jami, so'm", "s": "head"}],
            [{"v": 1, "f": "int"}, {"v": "Laminat"}, {"v": 2383333, "f": "money"}],
            [{"v": 2, "f": "int"}, {"v": "=HYPERLINK(\"http://x\")"}, {"v": 12.5, "f": "dec2"}],
            None,
            [None, {"v": "JAMI", "s": "grand"}, {"v": 2383346, "f": "money", "s": "grand"}],
        ]},
        {"name": "Xonalar", "cols": [5, 20], "rows": [[{"v": "Xonalar hisobi", "s": "title"}]]},
    ]}

    def test_excel_download(self):
        from io import BytesIO

        from openpyxl import load_workbook

        o = Obyekt.objects.create(owner=self.user, name="Chilonzor: 2/xona")
        r = self.c.post(reverse("obyekt_excel", args=[o.pk]), json.dumps(self.XLSX_SHEETS), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", r["Content-Disposition"])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Smeta", "Xonalar"])
        ws = wb["Smeta"]
        self.assertEqual(ws["A1"].value, "SMETA: Test")
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual(ws["B2"].fill.fgColor.rgb[-6:], "4D7C0F")  # sarlavha — lime (sayt palitrasi)
        self.assertEqual((ws["C3"].value, ws["C3"].number_format), (2383333, "#,##0"))
        self.assertEqual(ws["C4"].number_format, "0.00")
        self.assertEqual(ws["B4"].data_type, "s")  # "=..." formula emas, matn bo'lib qoladi
        self.assertEqual(ws["C6"].value, 2383346)
        self.assertEqual(ws.freeze_panes, "A3")
        self.assertEqual(ws.column_dimensions["B"].width, 30)
        self.assertIsNotNone(ws["A3"].border.left.style)  # jadval chegaralari

    def test_excel_rejects_bad_input_and_other_users(self):
        o = Obyekt.objects.create(owner=self.user)
        url = reverse("obyekt_excel", args=[o.pk])
        for bad in ("{bad", "{}", '{"sheets": []}', '{"sheets": [{"rows": [[{"v": true}]]}]}',
                    '{"sheets": [{"rows": [[{"v": 1e999}]]}]}'):
            self.assertEqual(self.c.post(url, bad, content_type="application/json").status_code, 400, bad)
        self.assertEqual(self.c.get(url).status_code, 405)
        other = Obyekt.objects.create(owner=User.objects.create_user("begona2", password="x"))
        r = self.c.post(reverse("obyekt_excel", args=[other.pk]), json.dumps(self.XLSX_SHEETS), content_type="application/json")
        self.assertEqual(r.status_code, 404)

    def test_landing_has_live_calculator(self):
        r = Client().get("/")
        html = r.content.decode()
        for needle in ('id="calc"', 'id="smeta-ref"', "smeta/js/calc.js", "smeta/js/landing.js",
                       'href="#qanday"', 'id="kim"', reverse("demo"), reverse("help"), reverse("privacy")):
            self.assertIn(needle, html)
        self.assertNotIn("smeta/js/app.js", html)  # og'ir ilova kodi bosh sahifada yuklanmaydi
        self.assertNotIn("PDF", html)  # PDF funksiyasi yo'q — va'da qilinmaydi

    def test_demo_page_and_public_excel(self):
        anon = Client()
        html = anon.get(reverse("demo")).content.decode()
        self.assertIn("SMETAGO_DEMO", html)
        self.assertNotIn('id="smeta-state"', html)  # namuna serverga yozilmaydi
        self.assertIn("smeta/js/calc.js", html)
        r = anon.post(reverse("demo_excel"), json.dumps(self.XLSX_SHEETS), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"PK"))  # xlsx = zip
        self.assertEqual(anon.post(reverse("demo_excel"), "{bad", content_type="application/json").status_code, 400)
        self.assertEqual(anon.get(reverse("demo_excel")).status_code, 405)
        # anonim endpoint: katta varaq (CPU) rad etiladi; kirgan egasi uchun chegara kengroq
        big = {"sheets": [{"name": "S", "cols": [10] * 8, "rows": [[{"v": 1}] * 8] * 1500}]}
        self.assertEqual(anon.post(reverse("demo_excel"), json.dumps(big), content_type="application/json").status_code, 400)
        huge = json.dumps({"sheets": [{"name": "S", "cols": [], "rows": [[{"v": "x" * 1000}]] * 400}]})
        self.assertEqual(anon.post(reverse("demo_excel"), huge, content_type="application/json").status_code, 413)
        o = Obyekt.objects.create(owner=self.user, state=STATE)
        r = self.c.post(reverse("obyekt_excel", args=[o.pk]), json.dumps(big), content_type="application/json")
        self.assertEqual(r.status_code, 200)
        strict = Client(enforce_csrf_checks=True)
        self.assertEqual(strict.post(reverse("demo_excel"), json.dumps(self.XLSX_SHEETS), content_type="application/json").status_code, 403)

    def test_help_and_privacy_pages(self):
        for name, uz, ru in (("help", "Qanday boshlayman?", "С чего начать?"), ("privacy", "Qanday ma'lumot saqlanadi", "Какие данные хранятся")):
            c = Client()
            self.assertContains(c.get(reverse(name)), uz)
            c.cookies["django_language"] = "ru"
            self.assertContains(c.get(reverse(name)), ru)

    def test_csrf_required_for_save(self):
        o = Obyekt.objects.create(owner=self.user)
        c = Client(enforce_csrf_checks=True)
        c.force_login(self.user)
        r = c.put(reverse("obyekt_state", args=[o.pk]), json.dumps(STATE), content_type="application/json")
        self.assertEqual(r.status_code, 403)

TG_TOKEN = "123456:TEST-token"


def tg_signed(**data):
    """Telegram Login Widget kabi imzolangan ma'lumot (sinov uchun)."""
    import hashlib
    import hmac
    import time

    data = {"auth_date": str(int(time.time())), **{k: str(v) for k, v in data.items()}}
    check = "\n".join(f"{k}={data[k]}" for k in sorted(data))
    data["hash"] = hmac.new(hashlib.sha256(TG_TOKEN.encode()).digest(), check.encode(), hashlib.sha256).hexdigest()
    return data


def tg_url(client, page="login"):
    """Sahifadagi widget qaytadigan manzil (sessiyaga bog'langan nonce bilan) — yo'l qismi."""
    import re

    html = client.get(reverse(page) if page != "/" else "/").content.decode()
    m = re.search(r'data-auth-url="https?://[^/"]+(/kirish/telegram/[^"]+)"', html)
    assert m, "widget topilmadi"
    return m.group(1)


TG_ON = {"TELEGRAM_BOT_TOKEN": TG_TOKEN, "TELEGRAM_BOT_NAME": "smetago_bot"}


class TelegramLoginTests(TestCase):
    def test_hidden_without_token(self):
        self.assertNotContains(Client().get(reverse("login")), "telegram-widget.js")
        self.assertEqual(Client().get(reverse("telegram_auth", args=["x"])).status_code, 404)

    def test_login_creates_account_then_reuses_it(self):
        from .models import TelegramAccount

        with self.settings(**TG_ON):
            c = Client()
            url = tg_url(c)
            self.assertContains(c.get(reverse("login")), 'data-telegram-login="smetago_bot"')
            r = c.get(url, tg_signed(id=777, first_name="Ali", username="ali_uz"))
            self.assertRedirects(r, "/")
            acct = TelegramAccount.objects.get(tg_id=777)
            self.assertEqual(acct.user.username, "ali_uz")
            self.assertFalse(acct.user.has_usable_password())
            self.assertContains(c.get("/"), 'id="kpis"')  # kirgan — dashboard
            # nonce bir martalik: shu manzilni qayta ishlatib bo'lmaydi
            self.assertEqual(Client().get(url, tg_signed(id=777, first_name="Ali")).status_code, 403)
            # ikkinchi marta (yangi sahifadan) — o'sha hisob, yangisi ochilmaydi
            c2 = Client()
            c2.get(tg_url(c2), tg_signed(id=777, first_name="Ali", username="ali_uz"))
            self.assertEqual(int(c2.session["_auth_user_id"]), acct.user.pk)
            self.assertEqual(User.objects.count(), 1)

    def test_forged_or_expired_rejected(self):
        with self.settings(**TG_ON):
            for bad in ({**tg_signed(id=5, first_name="X"), "id": "6"},   # boshqa odam nomidan
                        tg_signed(id=5, first_name="X", auth_date=1),     # eskirgan
                        {"id": "5"}):                                      # imzosiz
                c = Client()
                self.assertEqual(c.get(tg_url(c), bad).status_code, 403)
            self.assertEqual(User.objects.count(), 0)

    def test_link_to_existing_account(self):
        from .models import TelegramAccount

        u = User.objects.create_user("vali", password="SmetaGo-2026!")
        c = Client()
        c.force_login(u)
        with self.settings(**TG_ON):
            url = tg_url(c, "/")  # "Asosiy"dagi "Telegram'ni ulash"
            r = c.get(url, tg_signed(id=42, first_name="Vali"))
            self.assertRedirects(r, "/")
            self.assertEqual(TelegramAccount.objects.get(tg_id=42).user, u)
            self.assertNotContains(c.get("/"), "telegram-widget.js")  # ulangan — tugma yo'q
            anon = Client()
            anon.get(tg_url(anon), tg_signed(id=42, first_name="Vali"))
            self.assertEqual(int(anon.session["_auth_user_id"]), u.pk)

    def test_attacker_link_cannot_link_victim_account(self):
        """Hujum: hujumchi o'z sessiyasidagi manzil + o'z imzolangan ma'lumoti bilan havola yasaydi,
        kirgan qurbon uni ochadi — hujumchi Telegrami qurbon hisobiga ulanmasligi kerak."""
        from .models import TelegramAccount

        victim = User.objects.create_user("qurbon", password="SmetaGo-2026!")
        v = Client()
        v.force_login(victim)
        with self.settings(**TG_ON):
            attacker = Client()
            evil = tg_url(attacker)
            r = v.get(evil, tg_signed(id=666, first_name="Hujumchi"))
            self.assertEqual(r.status_code, 403)
            self.assertFalse(TelegramAccount.objects.filter(tg_id=666).exists())

    def test_attacker_link_cannot_log_in_victim(self):
        """Login CSRF: begona havola kirmagan qurbonni hujumchi hisobiga kiritmasligi kerak."""
        with self.settings(**TG_ON):
            attacker = Client()
            evil = tg_url(attacker)
            victim = Client()
            victim.get(reverse("login"))  # qurbonning o'z sessiyasi bor
            self.assertEqual(victim.get(evil, tg_signed(id=666, first_name="Hujumchi")).status_code, 403)
            self.assertNotIn("_auth_user_id", victim.session)

class ErrorPageTests(TestCase):
    def test_custom_404(self):
        r = Client().get("/bunday-sahifa-yoq/")
        self.assertEqual(r.status_code, 404)
        self.assertContains(r, "Sahifa topilmadi", status_code=404)
        self.assertContains(r, 'rel="manifest"', status_code=404)  # sayt uslubida (base.html)
